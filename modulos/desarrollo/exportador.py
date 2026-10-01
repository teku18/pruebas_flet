"""
Exportar Desarrollo a Excel: el reporte de todos los módulos con sus pendientes.

  Hoja "Resumen"
    ControlKraken · Desarrollo            (generado el …)
    Módulo | Descripción | Prioridad | Avance ▓▓▓░ | Pendientes | Alta | Media | Baja | Hechos | Total
    (cada módulo es un vínculo a su hoja; los conteos son FÓRMULAS que leen
     las hojas de cada módulo: si editas un pendiente en Excel, se actualizan)

  Una hoja por módulo ("Finanzas", "Agenda"…)
    Finanzas                                          ← Resumen
    Periodos, movimientos y reportes
    Prioridad Media · Avance 95 % · 3 pendientes · 1 hecho
    # | Prioridad | Estado | Tipo | Descripción / comentario | Adjuntos | Creado | Hecho
    (abiertos primero, de Alta a Baja; luego los hechos en gris)

Mismo estilo que el exportador de Finanzas: regresa bytes y la pantalla los
guarda con FilePicker.save_file (escritorio, web y celular).
"""
import math
import re
from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.formatting.rule import DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from modulos.desarrollo import services
from modulos.desarrollo.models import ModuloApp, Pendiente

# --- Formato -----------------------------------------------------------------
FUENTE = "Arial"
AZUL_OSCURO = "1F3864"
GRIS_TEXTO = "6B6A66"
LINEA = Side(style="thin", color="D0D0D0")
BORDE = Border(left=LINEA, right=LINEA, top=LINEA, bottom=LINEA)
FONDO_ENCABEZADO = PatternFill("solid", fgColor=AZUL_OSCURO)
FONDO_TOTAL = PatternFill("solid", fgColor="E8EAED")
FONDO_HECHO = PatternFill("solid", fgColor="F5F5F5")

# Prioridad: (relleno, color de letra) — mismo semáforo que la app
ESTILO_PRIORIDAD = {
    "alta": (PatternFill("solid", fgColor="FDE2E1"), "C00000"),
    "media": (PatternFill("solid", fgColor="FFF2CC"), "9C6500"),
    "baja": (PatternFill("solid", fgColor="E7ECEF"), "44546A"),
}

COLUMNAS_MODULO = ["#", "Prioridad", "Estado", "Tipo", "Descripción / comentario",
                   "Adjuntos", "Creado", "Hecho"]
ANCHOS_MODULO = [5, 11, 11, 13, 62, 30, 12, 12]
FILA_ENCABEZADO = 5            # encabezado de la tabla en cada hoja de módulo
MAX_FILA = 2000                # rango de las fórmulas del resumen

COLUMNAS_RESUMEN = ["Módulo", "Descripción", "Prioridad", "Avance", "Pendientes",
                    "Alta", "Media", "Baja", "Hechos", "Total"]
ANCHOS_RESUMEN = [18, 40, 11, 14, 12, 8, 8, 8, 9, 8]


def _font(**kw) -> Font:  # propio
    return Font(name=FUENTE, size=kw.pop("size", 10), **kw)


def _sheet_names(modulos: list[ModuloApp]) -> dict[int, str]:  # propio
    """Nombre de hoja válido y único por módulo (Excel: sin []:*?/\\, máx. 31)."""
    usados = {"resumen"}
    nombres = {}
    for m in modulos:
        base = re.sub(r"[\[\]:*?/\\]", "-", m.nombre).strip()[:28] or "Modulo"
        nombre, n = base, 2
        while nombre.lower() in usados:
            nombre = f"{base} {n}"
            n += 1
        usados.add(nombre.lower())
        nombres[m.id] = nombre
    return nombres


def _ref(hoja: str, rango: str) -> str:  # propio
    """'Mis cosas'!$C$6:$C$2000 (comillas siempre; una ' se escribe '')."""
    return f"'{hoja.replace(chr(39), chr(39) * 2)}'!{rango}"


def export_file_name() -> str:  # propio
    return f"Desarrollo ControlKraken {datetime.now():%Y-%m-%d}.xlsx"


# ============================================================================
# Entrada principal
# ============================================================================
def export_all() -> bytes:  # propio
    datos = services.all_for_export()
    nombres = _sheet_names([m for m, _ in datos])

    libro = Workbook()
    _summary_sheet(libro.active, datos, nombres)
    for m, pendientes in datos:
        _module_sheet(libro.create_sheet(nombres[m.id]), m, pendientes)

    salida = BytesIO()
    libro.save(salida)
    return salida.getvalue()


# ============================================================================
# Hoja Resumen
# ============================================================================
def _summary_sheet(hoja, datos, nombres: dict[int, str]):  # propio
    hoja.title = "Resumen"
    hoja.sheet_view.showGridLines = False

    hoja["A1"] = "ControlKraken · Desarrollo"
    hoja["A1"].font = _font(size=16, bold=True, color=AZUL_OSCURO)
    hoja["A2"] = f"Generado el {datetime.now():%d/%m/%Y %H:%M}"
    hoja["A2"].font = _font(color=GRIS_TEXTO, italic=True)

    fila_enc = 4
    for col, (titulo, ancho) in enumerate(zip(COLUMNAS_RESUMEN, ANCHOS_RESUMEN), start=1):
        celda = hoja.cell(fila_enc, col, titulo)
        celda.font = _font(bold=True, color="FFFFFF")
        celda.fill = FONDO_ENCABEZADO
        celda.alignment = Alignment(horizontal="center", vertical="center")
        celda.border = BORDE
        hoja.column_dimensions[get_column_letter(col)].width = ancho
    hoja.row_dimensions[fila_enc].height = 22

    primera = fila_enc + 1
    fila = primera
    for m, _ in datos:
        h = nombres[m.id]
        estado = _ref(h, f"$C${FILA_ENCABEZADO + 1}:$C${MAX_FILA}")
        prioridad = _ref(h, f"$B${FILA_ENCABEZADO + 1}:$B${MAX_FILA}")

        valores = [
            m.nombre,
            m.descripcion or "",
            m.prioridad_label,
            m.avance / 100,  # fracción: 0.95 se ve como 95 %
            f'=COUNTIF({estado},"Pendiente")',
            f'=COUNTIFS({estado},"Pendiente",{prioridad},"Alta")',
            f'=COUNTIFS({estado},"Pendiente",{prioridad},"Media")',
            f'=COUNTIFS({estado},"Pendiente",{prioridad},"Baja")',
            f'=COUNTIF({estado},"Hecho")',
            f"=E{fila}+I{fila}",
        ]
        for col, valor in enumerate(valores, start=1):
            celda = hoja.cell(fila, col, valor)
            celda.font = _font()
            celda.border = BORDE
            celda.alignment = Alignment(vertical="center", wrap_text=col == 2,
                                        horizontal="center" if col >= 3 else "left")

        # Módulo = vínculo a su hoja
        vinculo = hoja.cell(fila, 1)
        vinculo.hyperlink = f"#{_ref(h, 'A1')}"
        vinculo.font = _font(bold=True, color="0563C1", underline="single")

        relleno, color = ESTILO_PRIORIDAD[m.prioridad]
        hoja.cell(fila, 3).fill = relleno
        hoja.cell(fila, 3).font = _font(bold=True, color=color)
        hoja.cell(fila, 4).number_format = "0%"
        hoja.cell(fila, 6).font = _font(bold=True, color=ESTILO_PRIORIDAD["alta"][1])
        hoja.row_dimensions[fila].height = 30
        fila += 1

    ultima = fila - 1
    # Totales
    hoja.cell(fila, 1, "Total").font = _font(bold=True)
    for col in range(1, 11):
        celda = hoja.cell(fila, col)
        celda.fill = FONDO_TOTAL
        celda.border = BORDE
        celda.alignment = Alignment(horizontal="center" if col >= 3 else "left")
    if ultima >= primera:
        hoja.cell(fila, 4, f"=AVERAGE(D{primera}:D{ultima})").number_format = "0%"
        for col in range(5, 11):
            letra = get_column_letter(col)
            hoja.cell(fila, col, f"=SUM({letra}{primera}:{letra}{ultima})")
        for col in range(4, 11):
            hoja.cell(fila, col).font = _font(bold=True)
        # Barra de avance dentro de la celda (como un mini gráfico)
        hoja.conditional_formatting.add(
            f"D{primera}:D{ultima}",
            DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1,
                        color="5B9BD5"),
        )
    hoja.cell(fila + 1, 4, "Avance promedio").font = _font(size=8, color=GRIS_TEXTO)
    hoja.cell(fila + 1, 4).alignment = Alignment(horizontal="center")

    # Leyenda
    ley = fila + 3
    hoja.cell(ley, 1, "Prioridad:").font = _font(bold=True, color=GRIS_TEXTO)
    for i, clave in enumerate(("alta", "media", "baja")):
        relleno, color = ESTILO_PRIORIDAD[clave]
        celda = hoja.cell(ley + i, 2, clave.capitalize())
        celda.fill, celda.font = relleno, _font(bold=True, color=color)
    hoja.cell(ley, 3, "Toca el nombre de un módulo para ir a su hoja.").font = _font(
        italic=True, color=GRIS_TEXTO)

    hoja.freeze_panes = f"A{primera}"
    _print_setup(hoja)


# ============================================================================
# Hoja de un módulo
# ============================================================================
def _module_sheet(hoja, m: ModuloApp, pendientes: list[Pendiente]):  # propio
    hoja.sheet_view.showGridLines = False
    abiertos = sum(1 for p in pendientes if not p.hecho)
    hechos = len(pendientes) - abiertos

    hoja["A1"] = m.nombre
    hoja["A1"].font = _font(size=16, bold=True, color=AZUL_OSCURO)
    hoja["H1"] = "← Resumen"
    hoja["H1"].hyperlink = "#'Resumen'!A1"
    hoja["H1"].font = _font(color="0563C1", underline="single")
    hoja["A2"] = m.descripcion or ""
    hoja["A2"].font = _font(color=GRIS_TEXTO, italic=True)
    hoja["A3"] = (f"Prioridad {m.prioridad_label} · Avance {m.avance} % · "
                  f"{abiertos} pendiente{'s' if abiertos != 1 else ''} · "
                  f"{hechos} hecho{'s' if hechos != 1 else ''}")
    hoja["A3"].font = _font(bold=True, color=ESTILO_PRIORIDAD[m.prioridad][1])

    for col, (titulo, ancho) in enumerate(zip(COLUMNAS_MODULO, ANCHOS_MODULO), start=1):
        celda = hoja.cell(FILA_ENCABEZADO, col, titulo)
        celda.font = _font(bold=True, color="FFFFFF")
        celda.fill = FONDO_ENCABEZADO
        celda.alignment = Alignment(horizontal="center", vertical="center")
        celda.border = BORDE
        hoja.column_dimensions[get_column_letter(col)].width = ancho
    hoja.row_dimensions[FILA_ENCABEZADO].height = 22

    if not pendientes:
        hoja.cell(FILA_ENCABEZADO + 1, 5, "Sin pendientes 🎉").font = _font(
            italic=True, color=GRIS_TEXTO)

    for i, p in enumerate(pendientes, start=1):
        fila = FILA_ENCABEZADO + i
        adjuntos = "\n".join(a.nombre for a in p.adjuntos)
        valores = [
            i,
            p.prioridad_label,
            "Hecho" if p.hecho else "Pendiente",
            p.tipo_label,
            p.texto,
            adjuntos,
            p.creado_en.date() if p.creado_en else None,
            p.hecho_en.date() if p.hecho_en else None,
        ]
        gris = p.hecho
        for col, valor in enumerate(valores, start=1):
            celda = hoja.cell(fila, col, valor)
            celda.border = BORDE
            celda.font = _font(color=GRIS_TEXTO if gris else None)
            celda.alignment = Alignment(
                vertical="top", wrap_text=col in (5, 6),
                horizontal="left" if col in (5, 6) else "center",
            )
            if gris:
                celda.fill = FONDO_HECHO
        for col in (7, 8):
            hoja.cell(fila, col).number_format = "dd/mm/yyyy"

        relleno, color = ESTILO_PRIORIDAD[p.prioridad]
        if not p.hecho:
            hoja.cell(fila, 2).fill = relleno
            hoja.cell(fila, 2).font = _font(bold=True, color=color)
        hoja.cell(fila, 3).font = _font(bold=True, color="548235" if p.hecho else "C00000")

        # Alto del renglón según el texto (openpyxl no lo calcula solo)
        lineas_texto = sum(max(1, math.ceil(len(r) / 70)) for r in (p.texto or "").split("\n"))
        lineas = max(lineas_texto, len(p.adjuntos), 1)
        hoja.row_dimensions[fila].height = 15 * lineas + 4

    ultima = FILA_ENCABEZADO + max(len(pendientes), 1)
    hoja.auto_filter.ref = f"A{FILA_ENCABEZADO}:H{ultima}"
    hoja.freeze_panes = f"A{FILA_ENCABEZADO + 1}"
    hoja.print_title_rows = f"{FILA_ENCABEZADO}:{FILA_ENCABEZADO}"
    _print_setup(hoja)


def _print_setup(hoja):  # propio
    """Al imprimir: horizontal y todas las columnas en el ancho de la hoja."""
    hoja.page_setup.orientation = "landscape"
    hoja.page_setup.fitToWidth = 1
    hoja.page_setup.fitToHeight = 0
    hoja.sheet_properties.pageSetUpPr.fitToPage = True
