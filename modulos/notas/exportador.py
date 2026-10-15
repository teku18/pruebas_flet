"""
Exportar Notas a Excel (botón ▦ del tablero).

  Hoja "Notas"   todas (también las archivadas), una fila por nota:
                 Tipo | Título | Contenido | Etiqueta | Fijada | Archivada | Creada | Editada
                 (en las listas, Contenido = "☐ Leche · ☑ Pan · …")
  Una hoja por LISTA ("Súper", "Ideas 2"…): sus casillas con ☐ / ☑

Mismo estilo que los otros exportadores: regresa bytes y la pantalla los
guarda con FilePicker.save_file (compu y celular).
"""
import re
from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from modulos.notas.models import Nota

FUENTE = "Arial"
AZUL_OSCURO = "1F3864"
GRIS_TEXTO = "6B6A66"
LINEA = Side(style="thin", color="D0D0D0")
BORDE = Border(left=LINEA, right=LINEA, top=LINEA, bottom=LINEA)
FONDO_ENCABEZADO = PatternFill("solid", fgColor=AZUL_OSCURO)
FONDO_HECHO = PatternFill("solid", fgColor="F5F5F5")

COLUMNAS = ["Tipo", "Título", "Contenido", "Etiqueta", "Fijada", "Archivada", "Creada",
            "Editada"]
ANCHOS = [8, 26, 70, 14, 8, 10, 17, 17]


def _font(**kw) -> Font:  # propio
    return Font(name=FUENTE, size=kw.pop("size", 10), **kw)


def export_file_name() -> str:  # propio
    return f"Notas ControlKraken {datetime.now():%Y-%m-%d}.xlsx"


def _content(n: Nota) -> str:  # propio
    if n.es_lista:
        return " · ".join(("☑ " if i.hecho else "☐ ") + i.texto
                          for i in n.items if i.texto.strip())
    return n.cuerpo or ""


def _header(hoja, fila: int, titulos: list[str], anchos: list[int]):  # propio
    for col, (titulo, ancho) in enumerate(zip(titulos, anchos), start=1):
        celda = hoja.cell(fila, col, titulo)
        celda.font = _font(bold=True, color="FFFFFF")
        celda.fill = FONDO_ENCABEZADO
        celda.border = BORDE
        hoja.column_dimensions[celda.column_letter].width = ancho


def _sheet_names(listas: list[Nota]) -> dict[int, str]:  # propio
    """Nombre de hoja válido y único por lista (Excel: sin []:*?/\\, máx. 31)."""
    usados, nombres = {"notas"}, {}
    for n in listas:
        base = re.sub(r"[\[\]:*?/\\]", "-", n.titulo or "Lista").strip()[:28] or "Lista"
        nombre, k = base, 2
        while nombre.lower() in usados:
            nombre, k = f"{base} {k}", k + 1
        usados.add(nombre.lower())
        nombres[n.id] = nombre
    return nombres


def export_all() -> bytes:  # propio
    notas = Nota.search_all()  # fijadas primero, luego lo último editado
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Notas"
    hoja["A1"] = "ControlKraken · Notas"
    hoja["A1"].font = _font(size=14, bold=True)
    hoja["A2"] = f"Exportado el {datetime.now():%d/%m/%Y %H:%M} · {len(notas)} nota(s)"
    hoja["A2"].font = _font(color=GRIS_TEXTO)
    _header(hoja, 4, COLUMNAS, ANCHOS)
    hoja.freeze_panes = "A5"
    for fila, n in enumerate(notas, start=5):
        valores = ["Lista" if n.es_lista else "Nota", n.titulo or "", _content(n),
                   n.etiqueta.nombre if n.etiqueta else "", "Sí" if n.fijada else "",
                   "Sí" if n.archivada else "", n.creado_en, n.actualizado_en]
        for col, valor in enumerate(valores, start=1):
            celda = hoja.cell(fila, col, valor)
            celda.font = _font(color=GRIS_TEXTO if n.archivada else None)
            celda.border = BORDE
            celda.alignment = Alignment(vertical="top", wrap_text=col == 3)
            if col >= 7:
                celda.number_format = "dd/mm/yyyy hh:mm"

    listas = [n for n in notas if n.es_lista]
    nombres = _sheet_names(listas)
    for n in listas:
        h = libro.create_sheet(nombres[n.id])
        h["A1"] = n.titulo or "Lista"
        h["A1"].font = _font(size=14, bold=True)
        hechas, total = n.avance
        h["A2"] = f"✓ {hechas} de {total}" + (f" · {n.etiqueta.nombre}" if n.etiqueta else "")
        h["A2"].font = _font(color=GRIS_TEXTO)
        _header(h, 4, ["✓", "Elemento"], [5, 60])
        # Pendientes primero, como en la app
        for fila, i in enumerate(sorted(n.items, key=lambda x: (x.hecho, x.secuencia)), start=5):
            for col, valor in enumerate(["☑" if i.hecho else "☐", i.texto], start=1):
                celda = h.cell(fila, col, valor)
                celda.border = BORDE
                celda.font = _font(color=GRIS_TEXTO if i.hecho else None,
                                   strike=i.hecho and col == 2)
                if i.hecho:
                    celda.fill = FONDO_HECHO

    salida = BytesIO()
    libro.save(salida)
    return salida.getvalue()
