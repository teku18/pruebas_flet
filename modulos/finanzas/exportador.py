"""
Exportador a Excel (el camino de regreso del importador).

Genera un .xlsx de UN periodo (opcionalmente de un solo concepto) con dos hojas:

  1. <nombre del periodo>  -> los movimientos, con las MISMAS columnas que lee
                              el importador (más "Tipo"), así el archivo se
                              puede volver a importar tal cual.
       Inversion | Monto | Plataforma | Fecha | Entrada / Salida | Comentario | Tipo

  2. Reporte               -> lo mismo que la pantalla de Reportes:
       - Resumen: abre con, ingresos, gastos, cierra con (y meta/falta si es un concepto)
       - ¿Dónde está el dinero?  tabla por plataforma + gráfica de PASTEL
       - Ahorro mensual          tabla mes a mes + gráfica de LÍNEA del acumulado

Las gráficas son nativas de Excel y leen las celdas de las tablas: si cambias
un monto en la tabla, la gráfica se actualiza sola.

Los números por plataforma y por mes salen de services.py (mismas reglas que
la app: saldo inicial no es ingreso, traspasos internos se cancelan…). Los
totales, porcentajes y el acumulado son fórmulas de Excel.

Regresa bytes (no escribe a disco): Flet los guarda con FilePicker.save_file,
que así funciona igual en escritorio, web y celular.
"""
import re
from datetime import date
from io import BytesIO

from openpyxl import Workbook
from openpyxl.chart import LineChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.marker import Marker
from openpyxl.chart.series import DataPoint
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from core.ui import MESES, short_date
from modulos.finanzas.colores import AZUL, GRIS_OTRAS, platform_colors
from modulos.finanzas.models import TIPO_SELECTION, Inversion, Movimiento, Periodo
from modulos.finanzas.services import distribution_by_platform, monthly_savings, period_summary

# --- Formato -----------------------------------------------------------------
FUENTE = "Arial"
PESOS = '"$"#,##0.00;-"$"#,##0.00'                  # como tu Excel: $1,234.56
PESOS_SIGNO = '+"$"#,##0.00;-"$"#,##0.00;"$"0.00'   # del mes: +$5,996.92
PORCENTAJE = "0.00%"
FECHA = "dd/mm/yyyy"

# Rellenos de la hoja de movimientos (el azul es el de tus saldos iniciales)
FONDO_SALDO_INICIAL = PatternFill("solid", fgColor="6FA8DC")
FONDO_TRASPASO = PatternFill("solid", fgColor="FFF2CC")
FONDO_ENCABEZADO = PatternFill("solid", fgColor="E8EAED")
GRIS_TEXTO = "6B6A66"
LINEA = Side(style="thin", color="BFBFBF")

COLUMNAS = ["Inversion", "Monto", "Plataforma", "Fecha", "Entrada / Salida",
            "Comentario", "Tipo"]
ANCHOS = [16, 14, 16, 12, 16, 48, 18]


def _font(**kw) -> Font:  # propio
    return Font(name=FUENTE, size=kw.pop("size", 10), **kw)


def _sheet_name(texto: str) -> str:  # propio
    """Excel no acepta []:*?/\\ en el nombre de la hoja y máximo 31 caracteres."""
    return re.sub(r"[\[\]:*?/\\]", "-", texto).strip()[:31] or "Periodo"


def export_file_name(periodo: Periodo, concepto: Inversion | None = None) -> str:  # propio
    """'Inversiones 2026.xlsx' o 'Inversiones 2026 - Casa.xlsx'."""
    nombre = periodo.nombre + (f" - {concepto.nombre}" if concepto else "")
    return re.sub(r'[<>:"/\\|?*]', "-", nombre).strip() + ".xlsx"


# ============================================================================
# Entrada principal
# ============================================================================
def export_period(periodo: Periodo, inversion_id: int | None = None) -> bytes:  # propio
    """El .xlsx del periodo (con inversion_id: solo ese concepto) como bytes."""
    concepto = Inversion.get(inversion_id) if inversion_id else None
    libro = Workbook()
    _movements_sheet(libro.active, periodo, inversion_id)
    _report_sheet(libro.create_sheet("Reporte"), periodo, concepto)

    salida = BytesIO()
    libro.save(salida)
    return salida.getvalue()


# ============================================================================
# Hoja 1: movimientos (formato del importador)
# ============================================================================
def _movements_sheet(hoja, periodo: Periodo, inversion_id: int | None):  # propio
    hoja.title = _sheet_name(periodo.nombre)
    movimientos = [m for m in Movimiento.search_by_period(periodo.id)
                   if inversion_id is None or m.inversion_id == inversion_id]
    movimientos.sort(key=lambda m: (m.fecha, m.id))  # como en tu Excel: del más antiguo

    hoja.append(COLUMNAS)
    for celda in hoja[1]:
        celda.font = _font(bold=True)
        celda.fill = FONDO_ENCABEZADO

    for m in movimientos:
        hoja.append([
            m.inversion_label,
            m.monto,
            m.plataforma_label,
            m.fecha,
            "Entrada" if m.signo > 0 else "Salida",
            m.comentarios,
            # "Traspaso (sale)"/"(entra)" no es un Tipo del importador: al reimportar
            # se vuelven a emparejar solos (misma fecha y monto)
            TIPO_SELECTION.get(m.tipo, m.tipo),
        ])
        fila = hoja[hoja.max_row]
        for celda in fila:
            celda.font = _font()
        fila[1].number_format = PESOS
        fila[3].number_format = FECHA
        if m.tipo == "saldo_inicial":
            for celda in fila:
                celda.fill = FONDO_SALDO_INICIAL
        elif m.es_traspaso:
            for celda in fila:
                celda.fill = FONDO_TRASPASO

    for i, ancho in enumerate(ANCHOS, start=1):
        hoja.column_dimensions[get_column_letter(i)].width = ancho
    hoja.freeze_panes = "A2"
    if movimientos:
        hoja.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNAS))}{hoja.max_row}"


# ============================================================================
# Hoja 2: reporte con gráficas
# ============================================================================
#  Columnas: A = cuadrito de color | B = etiqueta | C = monto | D = % / acumulado
#            F en adelante = gráficas
def _report_sheet(hoja, periodo: Periodo, concepto: Inversion | None):  # propio
    hoja.sheet_view.showGridLines = False
    for col, ancho in {"A": 2.5, "B": 22, "C": 16, "D": 16, "E": 3}.items():
        hoja.column_dimensions[col].width = ancho

    titulo = concepto.nombre if concepto else "Todos los conceptos"
    hoja["B1"] = f"{periodo.nombre} · {titulo}"
    hoja["B1"].font = _font(size=14, bold=True)
    estado = " · cerrado" if periodo.cerrado else ""
    hoja["B2"] = (f"{short_date(periodo.fecha_inicio)} – {short_date(periodo.fecha_fin)}"
                  f"{estado} · exportado el {short_date(date.today())}")
    hoja["B2"].font = _font(color=GRIS_TEXTO)

    fila = _summary_block(hoja, 4, periodo, concepto)
    fila = _distribution_block(hoja, fila + 2, periodo, concepto)
    _monthly_block(hoja, fila + 2, periodo, concepto)


def _section_title(hoja, fila: int, titulo: str, subtitulo: str) -> int:  # propio
    hoja.cell(fila, 2, titulo).font = _font(size=12, bold=True)
    hoja.cell(fila + 1, 2, subtitulo).font = _font(size=9, color=GRIS_TEXTO)
    return fila + 2


def _table_header(hoja, fila: int, *titulos: str):  # propio
    for i, texto in enumerate(titulos):
        celda = hoja.cell(fila, 2 + i, texto)
        celda.font = _font(bold=True, color=GRIS_TEXTO)
        celda.border = Border(bottom=LINEA)
        if i:
            celda.alignment = Alignment(horizontal="right")


def _summary_block(hoja, fila: int, periodo: Periodo, concepto: Inversion | None) -> int:  # propio
    """Resumen del periodo. Regresa la última fila usada."""
    r = period_summary(periodo, concepto.id if concepto else None)
    fila = _section_title(hoja, fila, "Resumen", "Saldo inicial no cuenta como ingreso")
    inicio = fila
    renglones = [
        ("Abre con", r.saldo_apertura),
        ("Ingresos", r.ingresos),
        ("Gastos", r.gastos),
        ("Cierra con", f"=C{inicio}+C{inicio + 1}-C{inicio + 2}"),
    ]
    if concepto and concepto.meta:
        cierre = f"C{inicio + 3}"
        renglones += [
            ("Meta", concepto.meta),
            ("Falta", f"=MAX(C{inicio + 4}-{cierre},0)"),
            ("Avance", f"=IF(C{inicio + 4}>0,MIN({cierre}/C{inicio + 4},1),0)"),
        ]
    for i, (etiqueta, valor) in enumerate(renglones):
        f = inicio + i
        negrita = etiqueta == "Cierra con"
        hoja.cell(f, 2, etiqueta).font = _font(bold=negrita)
        celda = hoja.cell(f, 3, valor)
        celda.font = _font(bold=negrita)
        celda.number_format = PORCENTAJE if etiqueta == "Avance" else PESOS
        if negrita:
            hoja.cell(f, 2).border = hoja.cell(f, 3).border = Border(top=LINEA)
    return inicio + len(renglones) - 1


def _distribution_block(hoja, fila: int, periodo: Periodo,  # propio
                        concepto: Inversion | None) -> int:
    """Tabla por plataforma + pastel a la derecha. Regresa la última fila usada."""
    reparto = distribution_by_platform(periodo, concepto.id if concepto else None)
    fila = _section_title(hoja, fila, "¿Dónde está el dinero?",
                          "Saldo por plataforma al cierre del periodo")
    if not reparto:
        hoja.cell(fila, 2, "Sin saldo en plataformas para este filtro.").font = \
            _font(italic=True, color=GRIS_TEXTO)
        return fila

    colores = platform_colors()
    encabezado = fila
    _table_header(hoja, encabezado, "Plataforma", "Monto", "%")
    primera = encabezado + 1
    ultima = primera + len(reparto) - 1
    total = ultima + 1
    for i, p in enumerate(reparto):  # de menor a mayor, como en la app y tu Excel
        f = primera + i
        color = colores.get(p.plataforma_id, GRIS_OTRAS).lstrip("#").upper()
        hoja.cell(f, 1).fill = PatternFill("solid", fgColor=color)  # cuadrito de color
        hoja.cell(f, 2, p.nombre).font = _font()
        hoja.cell(f, 3, p.monto).number_format = PESOS
        hoja.cell(f, 4, f"=IF($C${total}=0,0,C{f}/$C${total})").number_format = PORCENTAJE
        hoja.cell(f, 3).font = hoja.cell(f, 4).font = _font()
    hoja.cell(total, 2, "Total")
    hoja.cell(total, 3, f"=SUM(C{primera}:C{ultima})").number_format = PESOS
    hoja.cell(total, 4, f"=SUM(D{primera}:D{ultima})").number_format = PORCENTAJE
    for col in (2, 3, 4):
        hoja.cell(total, col).font = _font(bold=True)
        hoja.cell(total, col).border = Border(top=LINEA)

    # --- Pastel: lee Plataforma (B) y Monto (C), sin el total ---------------
    pastel = PieChart()
    pastel.title = "Reparto por plataforma"
    pastel.add_data(Reference(hoja, min_col=3, min_row=encabezado, max_row=ultima),
                    titles_from_data=True)
    pastel.set_categories(Reference(hoja, min_col=2, min_row=primera, max_row=ultima))
    serie = pastel.series[0]
    for i, p in enumerate(reparto):  # mismo color que en la app
        punto = DataPoint(idx=i)
        punto.graphicalProperties.solidFill = colores.get(p.plataforma_id, GRIS_OTRAS).lstrip("#")
        punto.graphicalProperties.line.solidFill = "FFFFFF"  # separación blanca
        serie.dPt.append(punto)
    serie.dLbls = DataLabelList()
    serie.dLbls.showPercent = True
    serie.dLbls.showVal = serie.dLbls.showCatName = serie.dLbls.showSerName = False
    serie.dLbls.showLeaderLines = True
    serie.dLbls.position = "outEnd"   # % por fuera de la rebanada, como en la app
    serie.dLbls.numFmt = "0.00%"
    pastel.legend.position = "r"
    pastel.width, pastel.height = 14, 8   # cm
    hoja.add_chart(pastel, f"F{encabezado - 2}")

    # El pastel ocupa ~16 filas: la siguiente sección empieza debajo de él
    return max(total, encabezado + 14)


def _monthly_block(hoja, fila: int, periodo: Periodo,  # propio
                   concepto: Inversion | None) -> int:
    """Tabla mes a mes + línea del acumulado a la derecha."""
    meses = monthly_savings(periodo, concepto.id if concepto else None)
    fila = _section_title(hoja, fila, "Ahorro mensual",
                          "Depósitos + rendimientos − retiros (sin saldo inicial ni traspasos)")
    if not meses:
        hoja.cell(fila, 2, "Sin movimientos en el periodo.").font = \
            _font(italic=True, color=GRIS_TEXTO)
        return fila

    encabezado = fila
    _table_header(hoja, encabezado, "Mes", "Del mes", "Acumulado")
    primera = encabezado + 1
    for i, m in enumerate(meses):
        f = primera + i
        hoja.cell(f, 2, f"{MESES[m.mes - 1]} {m.año}")
        hoja.cell(f, 3, m.neto).number_format = PESOS_SIGNO
        # Acumulado = el del mes anterior + lo de este mes (fórmula)
        hoja.cell(f, 4, f"=C{f}" if i == 0 else f"=D{f - 1}+C{f}").number_format = PESOS
        for col in (2, 3, 4):
            hoja.cell(f, col).font = _font()
    ultima = primera + len(meses) - 1

    # --- Línea del acumulado --------------------------------------------------
    linea = LineChart()
    linea.title = "Acumulado del periodo"
    linea.add_data(Reference(hoja, min_col=4, min_row=encabezado, max_row=ultima),
                   titles_from_data=True)
    linea.set_categories(Reference(hoja, min_col=2, min_row=primera, max_row=ultima))
    serie = linea.series[0]
    serie.graphicalProperties.line.solidFill = AZUL.lstrip("#")
    serie.graphicalProperties.line.width = 28575  # 2.25 pt
    serie.marker = Marker(symbol="circle", size=6)
    serie.marker.graphicalProperties.solidFill = AZUL.lstrip("#")
    serie.marker.graphicalProperties.line.solidFill = AZUL.lstrip("#")
    serie.smooth = False
    linea.y_axis.number_format = '"$"#,##0'
    linea.y_axis.majorGridlines.spPr = None
    # openpyxl 3.1 oculta los ejes por defecto; así se ven en Excel y LibreOffice
    linea.x_axis.delete = False
    linea.x_axis.tickLblPos = "low"  # meses abajo aunque el acumulado sea negativo
    linea.y_axis.delete = False
    linea.legend = None  # una sola serie: el título basta
    linea.width, linea.height = 14, 8
    hoja.add_chart(linea, f"F{encabezado - 2}")
    return max(ultima, encabezado + 14)
