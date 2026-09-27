"""
Importador de Excel (como un wizard de importación de Odoo).

Formato esperado por pestaña (el de tu archivo):
  Inversion (o Concepto / Gasto) | Monto | Plataforma | Fecha | Entrada / Salida | Comentario

Reglas:
  - Cada pestaña es un Periodo (un "libro") con el nombre de la pestaña.
    Fechas: si el nombre trae un año ("Inversiones 2026") o todas las fechas
    son del mismo año, el periodo es ese año completo; si no, de la primera
    a la última fecha.
  - Si ya existe un periodo con ese nombre, la pestaña se omite (así no se
    duplica nada al reimportar). Para reimportarla, borra antes el periodo.
  - Plataformas y conceptos que no existan se crean en su catálogo.
  - Cada periodo es independiente: sus saldos iniciales son los del Excel.
  - Columna opcional "Tipo" (Saldo inicial / Depósito / Rendimiento / Retiro):
    si viene, manda. Si no, el tipo se deduce así:
      Bloque de Entradas al inicio de la hoja, del
        primer día, una por concepto/plataforma     -> saldo_inicial
        (termina en la primera Salida, otra fecha o un concepto/plataforma repetido)
      Comentario "saldo inicial"                    -> saldo_inicial
      Salida + Entrada misma fecha y monto, de un
        concepto/plataforma a otro distinto         -> traspaso (ligados)
      Entrada con "ganancia"/"rendimiento"          -> rendimiento
      Otra Entrada                                  -> deposito
      Otra Salida                                   -> retiro
Todo se guarda en una sola transacción: si algo falla, no se guarda nada.
"""
import re
import unicodedata
import uuid
from collections import Counter
from dataclasses import dataclass, field
from datetime import date, datetime

from openpyxl import load_workbook

from sqlalchemy import select

from core.database import SessionLocal
from modulos.finanzas.models import Inversion, Movimiento, Periodo, Plataforma

# Encabezado normalizado -> campo interno
COLUMNAS = {
    "inversion": "inversion",
    "concepto": "inversion",   # la columna puede llamarse Inversión, Concepto o Gasto
    "gasto": "inversion",
    "monto": "monto",
    "plataforma": "plataforma",
    "fecha": "fecha",
    "entrada / salida": "direccion",
    "entrada/salida": "direccion",
    "tipo": "tipo_excel",      # opcional: Saldo inicial / Depósito / Rendimiento / Retiro
    "comentario": "comentarios",
    "comentarios": "comentarios",
}
OBLIGATORIAS = {"inversion", "monto", "plataforma", "fecha", "direccion"}
PALABRAS_RENDIMIENTO = ("ganancia", "rendimiento", "interes")

# Columna opcional "Tipo": si viene, manda sobre las reglas automáticas
TIPOS_EXCEL = {
    "saldo inicial": "saldo_inicial",
    "deposito": "deposito",
    "rendimiento": "rendimiento",
    "retiro": "retiro",
}


@dataclass
class ImportResult:
    periodos: list[str] = field(default_factory=list)
    movimientos: int = 0
    por_tipo: Counter = field(default_factory=Counter)
    catalogos_nuevos: list[str] = field(default_factory=list)
    omitidas: list[str] = field(default_factory=list)   # "hoja: motivo"
    filas_con_error: list[str] = field(default_factory=list)
    avisos: list[str] = field(default_factory=list)       # posibles duplicados
    notas: list[str] = field(default_factory=list)        # decisiones automáticas


def _normalize(texto) -> str:  # propio
    """'Inversión ' -> 'inversion' (sin acentos, minúsculas, sin espacios extra)."""
    texto = unicodedata.normalize("NFKD", str(texto or "")).encode("ascii", "ignore").decode()
    return " ".join(texto.lower().split())


def _to_date(valor) -> date | None:  # propio
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    if isinstance(valor, str):
        for formato in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
            try:
                return datetime.strptime(valor.strip(), formato).date()
            except ValueError:
                pass
    return None


def _read_sheet(hoja) -> tuple[list[dict], list[str]]:  # propio
    """Filas de la hoja como dicts con los campos internos, y errores por fila."""
    filas_iter = hoja.iter_rows(values_only=True)
    encabezado = next(filas_iter, None) or []
    indices = {}
    for i, titulo in enumerate(encabezado):
        campo = COLUMNAS.get(_normalize(titulo))
        if campo:
            indices[campo] = i
    faltan = OBLIGATORIAS - set(indices)
    if faltan:
        raise ValueError(f"faltan columnas: {', '.join(sorted(faltan))}")

    filas, errores = [], []
    for num, valores in enumerate(filas_iter, start=2):
        if not valores or all(v is None for v in valores):
            continue
        get = lambda c: valores[indices[c]] if c in indices and indices[c] < len(valores) else None  # noqa: E731
        fecha = _to_date(get("fecha"))
        direccion = _normalize(get("direccion"))
        try:
            monto = abs(float(get("monto")))
        except (TypeError, ValueError):
            monto = None
        if fecha is None and monto is None and not get("inversion"):
            continue  # renglón plantilla (p. ej. solo trae "Nu" y "Entrada"): se ignora
        if not (fecha and monto and get("inversion") and get("plataforma")
                and direccion in ("entrada", "salida")):
            errores.append(f"{hoja.title} fila {num}: datos incompletos")
            continue
        filas.append({
            "fila": num,
            "inversion": str(get("inversion")).strip(),
            "plataforma": str(get("plataforma")).strip(),
            "monto": round(monto, 2),
            "fecha": fecha,
            "direccion": direccion,
            "comentarios": (str(get("comentarios")).strip() or None) if get("comentarios") else None,
            "tipo_excel": TIPOS_EXCEL.get(_normalize(get("tipo_excel"))),
        })
    return filas, errores


def _period_dates(nombre: str, filas: list[dict]) -> tuple[date, date]:  # propio
    """Año en el nombre (o todas las fechas del mismo año) -> año completo."""
    fechas = [f["fecha"] for f in filas]
    año = re.search(r"(19|20)\d{2}", nombre)
    años = {f.year for f in fechas}
    if año or len(años) == 1:
        a = int(año.group()) if año else años.pop()
        inicio, fin = date(a, 1, 1), date(a, 12, 31)
        return min([inicio, *fechas]), max([fin, *fechas])
    return min(fechas), max(fechas)


def _key(f: dict) -> tuple[str, str]:  # propio
    return _normalize(f["inversion"]), _normalize(f["plataforma"])


def _classify(filas: list[dict]) -> None:  # propio
    """Pone 'tipo' (y 'grupo' en traspasos) a cada fila. Modifica la lista."""
    # 1. Saldos iniciales: el bloque de Entradas con que empieza la hoja
    #    (mismo día, una por concepto/plataforma). Ej. Gastos: Pension, Renta,
    #    Vacaciones, Mio... y en cuanto aparece "Mio - Pago tarjeta" (Salida), termina.
    if filas and not filas[0]["tipo_excel"]:
        dia, vistos = filas[0]["fecha"], set()
        for f in filas:
            if (f["direccion"] != "entrada" or f["fecha"] != dia
                    or f["tipo_excel"] or _key(f) in vistos):
                break
            vistos.add(_key(f))
            f["tipo"] = "saldo_inicial"

    # 2. Traspasos: cada Salida busca su Entrada gemela (misma fecha y monto,
    #    distinto concepto y/o plataforma). Si hay varias, la fila más cercana.
    #    Las filas con Tipo explícito no se emparejan: ya dicen qué son.
    libre = lambda f: "tipo" not in f and not f["tipo_excel"]  # noqa: E731
    entradas_libres = [f for f in filas if f["direccion"] == "entrada" and libre(f)]
    for salida in [f for f in filas if f["direccion"] == "salida" and libre(f)]:
        candidatas = [e for e in entradas_libres
                      if e["fecha"] == salida["fecha"] and e["monto"] == salida["monto"]
                      and _key(e) != _key(salida)]
        if not candidatas:
            continue
        entrada = min(candidatas, key=lambda e: abs(e["fila"] - salida["fila"]))
        grupo = uuid.uuid4().hex
        salida.update(tipo="traspaso_salida", grupo=grupo)
        entrada.update(tipo="traspaso_entrada", grupo=grupo)
        # la entrada hereda el comentario de la salida si no trae uno
        entrada["comentarios"] = entrada["comentarios"] or salida["comentarios"]
        entradas_libres.remove(entrada)

    # 3. El resto
    for f in filas:
        if "tipo" in f:
            continue
        comentario = _normalize(f["comentarios"])
        if f["tipo_excel"]:
            f["tipo"] = f["tipo_excel"]  # lo que diga la columna Tipo
        elif f["direccion"] == "salida":
            f["tipo"] = "retiro"
        elif "saldo inicial" in comentario:
            f["tipo"] = "saldo_inicial"
        elif any(p in comentario for p in PALABRAS_RENDIMIENTO):
            f["tipo"] = "rendimiento"
        else:
            f["tipo"] = "deposito"


def import_workbook(ruta: str) -> ImportResult:  # propio
    """Importa todas las pestañas del archivo. Regresa un resumen."""
    resultado = ImportResult()
    libro = load_workbook(ruta, data_only=True, read_only=True)

    with SessionLocal() as session:
        catalogo_cache: dict[tuple, int] = {}

        def catalog_id(modelo, nombre: str) -> int:  # propio
            clave = (modelo.__name__, _normalize(nombre))
            if clave not in catalogo_cache:
                existia = modelo.find_by_name(nombre, session) is not None
                registro = modelo.get_or_create(nombre, session)
                if not existia:
                    resultado.catalogos_nuevos.append(f"{modelo.__name__}: {registro.nombre}")
                catalogo_cache[clave] = registro.id
            return catalogo_cache[clave]

        # 1. Leer todas las pestañas válidas
        hojas = []
        for hoja in libro.worksheets:
            nombre = hoja.title.strip()
            try:
                filas, errores = _read_sheet(hoja)
            except ValueError as ex:
                resultado.omitidas.append(f"{nombre}: {ex}")
                continue
            resultado.filas_con_error += errores
            if not filas:
                resultado.omitidas.append(f"{nombre}: sin filas válidas")
                continue
            if session.query(Periodo).filter(Periodo.nombre == nombre).first():
                resultado.omitidas.append(f"{nombre}: ya existe un periodo con ese nombre")
                continue
            hojas.append((nombre, filas, *_period_dates(nombre, filas)))

        # 2. Crear cada periodo con sus movimientos
        for nombre, filas, inicio, fin in sorted(hojas, key=lambda h: h[2]):
            periodo = Periodo(nombre=nombre, fecha_inicio=inicio, fecha_fin=fin, cerrado=False)
            session.add(periodo)
            session.flush()  # para tener periodo.id

            _classify(filas)
            saldos = [f for f in filas if f["tipo"] == "saldo_inicial"]
            if saldos:
                resultado.notas.append(
                    f"{nombre}: abre con {len(saldos)} saldo(s) inicial(es) "
                    f"(${sum(f['monto'] for f in saldos):,.2f}); no cuentan como ingreso."
                )
            entre_conceptos = [f for f in filas if f["tipo"] == "traspaso_salida"
                               and any(e.get("grupo") == f["grupo"]
                                       and _normalize(e["inversion"]) != _normalize(f["inversion"])
                                       for e in filas if e is not f)]
            if entre_conceptos:
                resultado.notas.append(
                    f"{nombre}: {len(entre_conceptos)} traspaso(s) entre conceptos detectados "
                    "(salida y entrada del mismo monto el mismo día). Revísalos: si alguno "
                    "no era traspaso, bórralo y captúralo como retiro + depósito."
                )

            # ¿Otro periodo ABIERTO con los mismos conceptos en estas fechas?
            # Puede ser el mismo dinero importado dos veces (p. ej. muestra y reales)
            conceptos = {_normalize(f["inversion"]) for f in filas}
            encimados = session.execute(
                select(Periodo.nombre, Inversion.nombre)
                .join(Movimiento, Movimiento.periodo_id == Periodo.id)
                .join(Inversion, Movimiento.inversion_id == Inversion.id)
                .where(Periodo.id != periodo.id, Periodo.cerrado.is_(False),
                       Movimiento.fecha >= inicio, Movimiento.fecha <= fin)
                .distinct()
            ).all()
            chocan = sorted({p for p, inv in encimados if _normalize(inv) in conceptos})
            if chocan:
                resultado.avisos.append(
                    f"{nombre}: «{', '.join(chocan)}» ya tiene movimientos de los mismos "
                    "conceptos en esas fechas. Si es el mismo dinero, se contará doble "
                    "en el acumulado global."
                )

            for f in filas:
                session.add(Movimiento(
                    periodo_id=periodo.id,
                    inversion_id=catalog_id(Inversion, f["inversion"]),
                    plataforma_id=catalog_id(Plataforma, f["plataforma"]),
                    monto=f["monto"],
                    fecha=f["fecha"],
                    tipo=f["tipo"],
                    comentarios=f["comentarios"],
                    traspaso_grupo=f.get("grupo"),
                ))
                resultado.por_tipo[f["tipo"]] += 1
            resultado.movimientos += len(filas)
            resultado.periodos.append(nombre)

        session.commit()  # todo o nada
    libro.close()
    return resultado
